from __future__ import annotations
import json, os, shutil, subprocess, sys, tempfile, hashlib
from functools import lru_cache
from pathlib import Path
from typing import Any
import pymupdf as fitz
from bs4 import BeautifulSoup
from docx import Document
from .cleaning import clean_text, remove_repeated_page_lines
from .util import write_jsonl, stable_id


def extract_pdf_native(path: Path) -> tuple[str, list[str]]:
    doc = fitz.open(path)
    pages = [page.get_text("text") or "" for page in doc]
    return "\n\n".join(pages), pages


@lru_cache(maxsize=2)
def _docling_converter(ocr_languages: tuple[str,...], use_gpu: bool):
    try:
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions, EasyOcrOptions
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.backend.pypdfium2_backend import PyPdfiumDocumentBackend
    except Exception as e:
        raise RuntimeError("Docling/OCR extras are not installed. Run RUN_0_SETUP_FULL.bat") from e
    opts = PdfPipelineOptions()
    opts.do_ocr = True
    opts.do_table_structure = False
    opts.ocr_options = EasyOcrOptions(lang=list(ocr_languages), use_gpu=use_gpu, force_full_page_ocr=True)
    # Docling's native parser cannot open its resources under accented Windows paths.
    converter = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(
        pipeline_options=opts, backend=PyPdfiumDocumentBackend)})
    return converter


def extract_pdf_docling(path: Path, ocr_languages: list[str], use_gpu: bool, page_number: int | None=None) -> str:
    converter=_docling_converter(tuple(ocr_languages),use_gpu)
    result=converter.convert(path,**({'page_range':(page_number,page_number)} if page_number else {}))
    return result.document.export_to_text()


def extract_pdf_pages(path: Path, row: dict, cfg: dict, output_dir: Path):
    settings=cfg['extraction']; mode=str(settings.get('use_docling','auto')).lower()
    if mode not in {'auto','always','never'}: raise ValueError('Invalid use_docling mode')
    threshold=int(settings.get('min_text_chars_before_ocr',1200)); pages=[]; provenance=[]
    cache=output_dir/'01_extracted'/'page_cache'; cache.mkdir(parents=True,exist_ok=True)
    with fitz.open(path) as pdf:
        for i,page in enumerate(pdf):
            native=page.get_text('text') or ''; count=len(''.join(native.split()))
            images=page.get_image_info()
            image_area=max((fitz.Rect(image['bbox']).get_area() for image in images),default=0)
            coverage=image_area/max(page.rect.get_area(),1)
            blank=count==0 and not images and not page.get_drawings()
            bad_fraction=(native.count('\ufffd')+native.count('\x00'))/max(len(native),1)
            needs=not blank and (mode=='always' or count<40 or
                (count<threshold and coverage>.25) or (coverage>.6 and count<2000) or bad_fraction>.02)
            text=native; method='pymupdf'; error=None; status='NOT_NEEDED'
            key=stable_id('page-v3',row['sha256'],str(i+1),json.dumps(settings,sort_keys=True),prefix='page')
            cached=cache/(key+'.json')
            if needs and mode!='never':
                try:
                    if cached.exists():
                        text=json.loads(cached.read_text(encoding='utf-8'))['text']; method='docling_easyocr_cached'
                    else:
                        text=extract_pdf_docling(path,settings.get('ocr_languages',['vi','en']),bool(settings.get('ocr_use_gpu',False)),i+1)
                        if not text.strip(): raise RuntimeError('OCR returned empty text')
                        temporary=cached.with_suffix('.tmp')
                        temporary.write_text(json.dumps({'text':text},ensure_ascii=False),encoding='utf-8'); os.replace(temporary,cached)
                        method='docling_easyocr'
                    status='COMPLETE'
                except Exception as exc:
                    error=str(exc); status='FAILED'; text=native
            elif needs: status='DISABLED'
            pages.append(text)
            provenance.append({'page':i+1,'method':method,'chars':len(text),'native_chars':count,
                'image_coverage':round(coverage,3),'bad_character_fraction':bad_fraction,'ocr':method.startswith('docling'),'ocr_status':status,
                'blank':blank,'error':error,'file_id':row['file_id'],'sha256':row['sha256']})
    return pages,provenance


def extract_docx(path: Path) -> str:
    doc = Document(path)
    blocks: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip(): blocks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def convert_legacy_doc(path: Path, cache_dir: Path) -> Path | None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    out = cache_dir / (path.stem + ".docx")
    if out.exists() and out.stat().st_size > 100:
        return out
    # Microsoft Word COM is the most reliable on the user's Windows environment.
    if sys.platform == "win32":
        try:
            import win32com.client  # type: ignore
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(str(path.resolve()))
            doc.SaveAs2(str(out.resolve()), FileFormat=16)  # wdFormatDocumentDefault (.docx)
            doc.Close(False)
            word.Quit()
            if out.exists(): return out
        except Exception:
            try:
                word.Quit()  # type: ignore[name-defined]
            except Exception:
                pass
    # LibreOffice fallback.
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", str(cache_dir), str(path)],
                           check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
            if out.exists(): return out
        except Exception:
            pass
    # antiword fallback gives plain text; caller handles separately.
    return None


def extract_doc(path: Path, cache_dir: Path) -> tuple[str, str]:
    converted = convert_legacy_doc(path, cache_dir)
    if converted:
        return extract_docx(converted), "word_or_libreoffice_conversion"
    antiword = shutil.which("antiword")
    if antiword:
        proc = subprocess.run([antiword, str(path)], capture_output=True, timeout=120)
        for enc in ("utf-8", "cp1258", "latin-1"):
            try:
                return proc.stdout.decode(enc), "antiword"
            except Exception:
                continue
    raise RuntimeError("Cannot read legacy .doc. Install Microsoft Word, LibreOffice, or antiword.")


def extract_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer"]):
        tag.decompose()
    return soup.get_text("\n")


def extract_one(row: dict, cfg: dict, output_dir: Path) -> dict:
    path = Path(row["path"])
    ext = path.suffix.lower()
    method = ""
    pages: list[str] = []
    needs_ocr = False
    error = None
    raw = ""
    page_provenance=[]
    try:
        if ext == ".pdf":
            pages,page_provenance=extract_pdf_pages(path,row,cfg,output_dir)
            method='hybrid_page_pdf'; needs_ocr=any(p['ocr_status'] in {'FAILED','DISABLED'} for p in page_provenance)
            if needs_ocr: error='One or more pages require OCR; see page_provenance.'
        elif ext == ".docx":
            raw = extract_docx(path); pages = [raw]; method = "python-docx"
        elif ext == ".doc":
            raw, method = extract_doc(path, output_dir / "01_extracted" / "converted_docx"); pages = [raw]
        elif ext in {".html", ".htm"}:
            raw = extract_html(path); pages = [raw]; method = "beautifulsoup"
        elif ext == ".txt":
            raw = path.read_text(encoding="utf-8", errors="ignore"); pages = [raw]; method = "text"
        elif ext == ".json":
            raw = json.dumps(json.loads(path.read_text(encoding="utf-8")), ensure_ascii=False, indent=2); pages=[raw]; method="json"
        else:
            raise RuntimeError(f"Unsupported file type {ext}")
        if pages:
            pages = remove_repeated_page_lines(
                pages,
                float(cfg["cleaning"].get("repeated_line_page_ratio", .55)),
                int(cfg["cleaning"].get("repeated_line_max_chars", 120)),
            )
        cleaned_pages=[clean_text(p, cfg['cleaning'].get('unicode_form','NFC')) for p in pages]
        cleaned = '\n\n'.join(cleaned_pages) if pages else clean_text(raw,cfg['cleaning'].get('unicode_form','NFC'))
        offset=0
        for p,content in zip(page_provenance,cleaned_pages):
            p['text_start']=offset; p['text_end']=offset+len(content); p['cleaned_chars']=len(content); offset+=len(content)+2
    except Exception as e:
        cleaned = ""; error = str(e)
    out = {
        **row,
        "extraction_method": method,
        "text": cleaned,
        "text_chars": len(cleaned),
        "needs_ocr": needs_ocr,
        "extraction_error": error,
        'page_provenance':page_provenance,'page_count':len(page_provenance),'extraction_schema':'page-v3',
    }
    return out


def extract_all(manifest: list[dict], cfg: dict, output_dir: Path) -> list[dict]:
    from tqdm import tqdm
    rows=[]
    timeout=float(cfg['extraction'].get('document_timeout_seconds',180))
    if timeout<=0: raise ValueError('document_timeout_seconds must be positive')
    cache=output_dir/'01_extracted'/'document_cache'; cache.mkdir(parents=True,exist_ok=True)
    for row in tqdm(manifest,desc='Extract/Clean'):
        key=stable_id('document-v3',row['file_id'],row['sha256'],json.dumps({'extraction':cfg['extraction'],'cleaning':cfg['cleaning']},sort_keys=True),prefix='extract')
        cached=cache/(key+'.json')
        if cached.exists():
            result=json.loads(cached.read_text(encoding='utf-8'))
        else:
            with tempfile.TemporaryDirectory(dir=cache) as work:
                request=Path(work)/'request.json'; response=Path(work)/'response.json'
                request.write_text(json.dumps({'row':row,'cfg':cfg,'output_dir':str(output_dir),'response':str(response)},default=str),encoding='utf-8')
                try:
                    process=subprocess.run([sys.executable,'-m','vn_labor_offline.extraction_worker',str(request)],capture_output=True,timeout=timeout)
                    if process.returncode or not response.exists(): raise RuntimeError(process.stderr.decode('utf-8',errors='replace')[-1500:])
                    result=json.loads(response.read_text(encoding='utf-8'))
                except (subprocess.TimeoutExpired,RuntimeError) as exc:
                    result={**row,'text':'','text_chars':0,'needs_ocr':row['extension']=='.pdf','extraction_method':'failed_worker',
                        'page_provenance':[],'extraction_schema':'page-v3','extraction_error':f'{type(exc).__name__}: {exc}'}
            if not result.get('extraction_error') and not any(p.get('ocr_status')=='FAILED' for p in result.get('page_provenance',[])):
                cached.write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
        rows.append(result)
        write_jsonl(output_dir/'01_extracted'/'documents.jsonl',rows)
    write_jsonl(output_dir / "01_extracted" / "documents.jsonl", rows)
    failures = [r for r in rows if r.get("extraction_error") or r.get("text_chars", 0) < 100]
    write_jsonl(output_dir / "reports" / "extraction_issues.jsonl", failures)
    return rows
